#excel импорты
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter

#docx импорты
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

#импорты для Сценария 1
import csv
import json
from collections import defaultdict

#импорты для Сценария 3
import time
import tracemalloc
import random

class Uniimport:
    @staticmethod
    def read_range(file_path, sheet_name, cell_range):
        #Читает данные из указанного листа и диапазона ячеек.
        wb = load_workbook(file_path, data_only=True)

        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Лист '{sheet_name}' не найден. Доступные: {wb.sheetnames}")

        ws = wb[sheet_name]
        cells = ws[cell_range]  # например, A1:D10

        # Преобразуем в список списков
        data = [[cell.value for cell in row] for row in cells]

        wb.close()
        return data

    @staticmethod
    def print_structure(data):
        #Выводит структуру данных в консоль в виде таблицы.
        if not data:
            print("Данные отсутствуют")
            return

        # Приводим все значения к строкам для выравнивания
        str_data = [[("" if v is None else str(v)) for v in row] for row in data]

        # Определяем ширину каждой колонки
        n_cols = max(len(row) for row in str_data)
        widths = [0] * n_cols
        for row in str_data:
            for i, val in enumerate(row):
                widths[i] = max(widths[i], len(val))


        for i, row in enumerate(str_data):
            line = "|" + "|".join(f" {val:<{widths[j]}} " for j, val in enumerate(row)) + "|"
            print(line)


    @staticmethod
    def create_report(source_file, source_sheet, cell_range, report_file="report_module1.xlsx"):
        #Создаёт отчёт на основе данных из источника с форматированием.
        #Читаем исходные данные
        data = Uniimport.read_range(source_file, source_sheet, cell_range)

        #Создаём новый файл отчёта
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        #Записываем данные
        for row in data:
            ws.append(row)

        #Жирный шрифт для заголовков (первая строка)
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        #Границы для всех ячеек
        thin = Side(border_style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        #Выравнивание
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        n_rows = len(data)
        n_cols = max(len(row) for row in data) if data else 0

        for r in range(1, n_rows + 1):
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border

                if r == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center
                else:
                    cell.alignment = left if c == 2 else center

        #Авто ширина колонок
        for c in range(1, n_cols + 1):
            letter = get_column_letter(c)
            max_len = 0
            for r in range(1, n_rows + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[letter].width = max_len + 3

        #Закрепляем верхнюю строку
        ws.freeze_panes = "A2"

        #Сохраняем
        wb.save(report_file)
        print(f"Отчёт сохранён в файл: {report_file}")
        return report_file

    #Методы для docx
    @staticmethod
    def add_toc(doc):
        #Вставляет поле автоматического оглавления.
        p = doc.add_paragraph()
        run = p.add_run()

        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')

        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'TOC \\o "1-3" \\h \\z \\u'

        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')

        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    @staticmethod
    def add_section(doc, title, level, text):
        """Добавляет раздел с заголовком и абзацем."""
        doc.add_heading(title, level=level)
        doc.add_paragraph(text)

    @staticmethod
    def add_numbered_list(doc, items):
        """Добавляет нумерованный список."""
        for item in items:
            doc.add_paragraph(item, style='List Number')

    @staticmethod
    def create_document(path="document.docx"):
        """Создаёт документ с оглавлением, разделами и списком."""
        doc = Document()

        # Оглавление
        doc.add_heading("Содержание", level=1)
        Uniimport.add_toc(doc)
        doc.add_page_break()

        # Разделы с разными уровнями заголовков
        Uniimport.add_section(doc, "Введение", 1,
                    "Краткое введение в документ.")
        Uniimport.add_section(doc, "Обзор", 2,
                    "Обзор рассматриваемой темы.")

        # Нумерованный список
        doc.add_heading("Список задач", level=2)
        Uniimport.add_numbered_list(doc, [
            "Первая задача",
            "Вторая задача",
            "Третья задача",
        ])

        Uniimport.add_section(doc, "Заключение", 1,
                    "Итоги и выводы.")

        doc.save(path)
        print(f"Сохранено: {path}")

    #Блок методов для модуля 3 (сценарии 1, 2, 3)

    #Сценарий 1

    #Чтение
    @staticmethod
    def read_csv(path):
        """Читает CSV-файл в список словарей."""
        with open(path, encoding="utf-8") as f:
            return list(csv.DictReader(f))

    @staticmethod
    def read_json(path):
        """Читает JSON-файл в список словарей."""
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else [data]

    #Агрегация
    @staticmethod
    def aggregate(rows, group_key, value_key):
        """
        Группирует строки по group_key и считает сумму, среднее
        и количество по value_key.
        """
        totals = defaultdict(float)
        counts = defaultdict(int)

        for row in rows:
            group = row[group_key]
            try:
                value = float(row[value_key])
            except (ValueError, TypeError):
                continue
            totals[group] += value
            counts[group] += 1

        result = []
        for group in sorted(totals):
            total = totals[group]
            count = counts[group]
            result.append({
                group_key: group,
                "Сумма": round(total, 2),
                "Среднее": round(total / count, 2) if count else 0,
                "Количество": count,
            })
        return result

    #Экспорт
    @staticmethod
    def export_aggregate_to_excel(aggregated, path, group_key):
        """Сохраняет агрегированную таблицу в Excel с форматированием."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        headers = [group_key, "Сумма", "Среднее", "Количество"]
        ws.append(headers)

        for row in aggregated:
            ws.append([row[h] for h in headers])

        # Форматирование
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", start_color="4472C4")
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
            for cell in row:
                cell.border = border

        ws.freeze_panes = "A2"

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

        wb.save(path)

    @staticmethod
    def export_aggregate_to_word(aggregated, path, group_key, title="Отчёт по продажам"):
        """Сохраняет агрегированную таблицу в Word как отчёт."""
        doc = Document()

        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        headers = [group_key, "Сумма", "Среднее", "Количество"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"

        # Шапка
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(11)

        # Данные
        for row in aggregated:
            cells = table.add_row().cells
            for i, h in enumerate(headers):
                cells[i].text = str(row[h])

        doc.save(path)

    #Главная функция Сценария 1
    @staticmethod
    def build_report(source_path, group_key, value_key,
                     excel_path="report.xlsx", word_path="report.docx"):
        """Полный конвейер: чтение → агрегация → экспорт."""
        if source_path.lower().endswith(".json"):
            rows = Uniimport.read_json(source_path)
        else:
            rows = Uniimport.read_csv(source_path)

        aggregated = Uniimport.aggregate(rows, group_key, value_key)

        Uniimport.export_aggregate_to_excel(aggregated, excel_path, group_key)
        Uniimport.export_aggregate_to_word(aggregated, word_path, group_key)

        return aggregated

    #Сценарий 2
    @staticmethod
    def read_word_table(path, table_index=0, header_row=0):
        """
        Читает таблицу из Word-документа и возвращает список словарей.
        """
        doc = Document(path)

        if table_index >= len(doc.tables):
            raise ValueError(
                f"В документе {len(doc.tables)} таблиц, "
                f"запрошена #{table_index}"
            )

        table = doc.tables[table_index]

        # Заголовки
        headers = [cell.text.strip() for cell in table.rows[header_row].cells]

        # Данные
        records = []
        for row in table.rows[header_row + 1:]:
            values = [cell.text.strip() for cell in row.cells]
            records.append(dict(zip(headers, values)))

        return records

    @staticmethod
    def records_to_objects(records, cls):
        """
        Преобразует список словарей в список объектов указанного класса.
        Поля класса должны совпадать с именами колонок (или приводиться).
        """
        objects = []
        for rec in records:
            try:
                obj = cls(**rec)
            except TypeError:
                # Пробуем привести ключи к нижнему регистру
                obj = cls(**{k.lower(): v for k, v in rec.items()})
            objects.append(obj)
        return objects

    #Сценарий 3
    @staticmethod
    def generate_dataset(rows=50_000):
        """Генерирует список строк со случайными данными."""
        categories = ["Книги", "Игры", "Одежда", "Продукты", "Техника"]
        data = []
        for i in range(1, rows + 1):
            data.append([
                i,
                random.choice(categories),
                round(random.uniform(100, 10_000), 2),
                random.randint(1, 100),
            ])
        return data

    @staticmethod
    def export_to_excel(data, path="big_dataset.xlsx"):
        """Экспортирует данные в Excel. Возвращает время и память."""
        tracemalloc.start()
        start = time.perf_counter()

        wb = Workbook(write_only=True)  # режим для больших файлов
        ws = wb.create_sheet("Data")

        ws.append(["ID", "Категория", "Цена", "Количество"])
        for row in data:
            ws.append(row)

        wb.save(path)

        elapsed = time.perf_counter() - start
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()

        return {
            "rows": len(data),
            "time_sec": round(elapsed, 2),
            "peak_memory_mb": round(peak / 1024 / 1024, 2),
            "file": path,
        }

    def benchmark_export(rows=50_000, path="big_dataset.xlsx"):
        """Полный замер: генерация + экспорт."""
        print(f"Генерация {rows} строк...")
        data = Uniimport.generate_dataset(rows)

        print("Экспорт в Excel...")
        stats = Uniimport.export_to_excel(data, path)

        print("\nРезультат")
        print(f"Строк:            {stats['rows']:,}")
        print(f"Время:            {stats['time_sec']} сек")
        print(f"Пик памяти:       {stats['peak_memory_mb']} МБ")
        print(f"Файл:             {stats['file']}")

        return stats

